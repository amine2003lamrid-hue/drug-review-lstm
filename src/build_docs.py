"""Fill report/report.md and slides/slides.md from results/, then export.

Every number comes from results/*.json, results/eda_stats.json or a model instantiation (parameter
counts). Anything that depends on a run that has not happened yet becomes **[TO FILL AFTER RUN]**.

Export: report -> PDF via pandoc if available, else .docx via python-docx.
        slides -> PDF + PPTX via marp-cli if available.
"""
import argparse
import glob
import json
import os
import re
import shutil
import subprocess

import yaml

from .cells import RecurrentLayer
from .make_report_assets import ARCH, BONUS, HPARAM, load
from .model import count_params

TODO = "**[TO FILL AFTER RUN]**"


def pct(x):
    return f"{100 * x:.1f}%"


def params_table(cfg, vocab_size):
    rows = ["| Cell | Recurrent params | Total params (incl. embedding + classifier) | vs vanilla (recurrent) |",
            "|---|---|---|---|"]
    D, H = cfg["embed_dim"], cfg["hidden_size"]
    base = None
    for name, cell, P in (("Vanilla LSTM", "vanilla", None), ("CIFG-LSTM", "cifg", None),
                          ("LSTMP (P=128)", "lstmp", 128), ("LSTMP (P=64)", "lstmp", 64)):
        rnn = RecurrentLayer(cell, D, H, proj_size=P or 128)
        n = count_params(rnn)
        total = n + vocab_size * D + (rnn.output_size + 1) * 2 if vocab_size else None
        base = base or n
        rows.append(f"| {name} | {n:,} | {f'{total:,}' if total else TODO} | {n / base - 1:+.1%} |")
    return "\n".join(rows)


def table_from(path, short=False):
    if not os.path.exists(path):
        return TODO
    with open(path) as f:
        lines = [l for l in f.read().strip().splitlines() if l.startswith("|")]
    if short:  # model | params | acc | F1 | time
        out = []
        for l in lines:
            c = [x.strip() for x in l.strip("|").split("|")]
            out.append("| " + " | ".join([c[0], c[1], c[3], c[4], c[8]]) + " |")
        return "\n".join(out)
    return "\n".join(lines)


def arch_findings(res):
    if "vanilla" not in res:
        return TODO, TODO
    v = res["vanilla"]
    out = []
    for k, lab in ARCH[1:]:
        if k not in res:
            out.append(f"{lab}: {TODO}")
            continue
        r = res[k]
        d = r["test"]["macro_f1"] - v["test"]["macro_f1"]
        out.append(f"{lab} reaches a test macro-F1 of {r['test']['macro_f1']:.4f} ({d:+.4f} vs vanilla) with "
                   f"{1 - r['recurrent_params'] / v['recurrent_params']:.1%} fewer recurrent parameters and "
                   f"{r['mean_epoch_time_s']:.1f} s/epoch vs {v['mean_epoch_time_s']:.1f} s for vanilla; "
                   f"its best validation epoch is {r['best_epoch']}.")
    long = ("The vanilla baseline reaches a test macro-F1 of {:.4f} (accuracy {:.4f}). ".format(
        v["test"]["macro_f1"], v["test"]["accuracy"]) + " ".join(out) +
        f" {TODO.replace('TO FILL AFTER RUN', 'TO FILL AFTER RUN: 1-2 sentences interpreting these gaps')}")
    return long, " ".join(out)


def hparam_findings(res):
    if "vanilla" not in res:
        return TODO, TODO
    b = res["vanilla"]["test"]["macro_f1"]
    parts, best = [], ("baseline", b)
    for k, lab, factor in HPARAM[1:]:
        if k not in res:
            parts.append(f"{lab}: {TODO}.")
            continue
        r = res[k]
        f1 = r["test"]["macro_f1"]
        if f1 > best[1]:
            best = (lab, f1)
        parts.append(f"*{factor}*: {lab} gives {f1:.4f} ({f1 - b:+.4f}), best epoch {r['best_epoch']}, "
                     f"{r['mean_epoch_time_s']:.1f} s/epoch.")
    long = (f"Relative to the baseline (macro-F1 {b:.4f}): " + " ".join(parts) +
            f" The best setting in this study is **{best[0]}** ({best[1]:.4f}). "
            f"{TODO.replace('TO FILL AFTER RUN', 'TO FILL AFTER RUN: interpret, e.g. SGD convergence speed, batch-size/time trade-off')}")
    short = f"Best: **{best[0]}** (macro-F1 {best[1]:.3f}) vs baseline {b:.3f}"
    return long, short


def proj_findings(res):
    a, b = (res.get(k) for k, _ in BONUS)
    if not (a and b):
        return f"**Projection size (bonus).** {TODO}"
    return (f"**Projection size (bonus).** Halving the projection from P=128 to P=64 reduces the recurrent "
            f"parameters from {a['recurrent_params']:,} to {b['recurrent_params']:,} "
            f"({1 - b['recurrent_params'] / a['recurrent_params']:.1%} fewer) and changes test macro-F1 from "
            f"{a['test']['macro_f1']:.4f} to {b['test']['macro_f1']:.4f} ({b['test']['macro_f1'] - a['test']['macro_f1']:+.4f}).")


def best_summary(res):
    if not res:
        return TODO, TODO, TODO
    best = max(res.values(), key=lambda r: r["best_val_macro_f1"])
    t = best["test"]
    (tn, fp), (fn, tp) = t["confusion_matrix"]
    s = (f"Selecting on **validation** macro-F1 across all runs, the final model is **{best['name']}** "
         f"(cell `{best['config']['cell']}`, {best['trainable_params']:,} trainable parameters, val macro-F1 "
         f"{best['best_val_macro_f1']:.4f}). On the {tn + fp + fn + tp:,}-review test set it obtains accuracy "
         f"**{t['accuracy']:.4f}**, macro-F1 **{t['macro_f1']:.4f}**, macro precision {t['precision']:.4f} and macro "
         f"recall {t['recall']:.4f}.")
    cm = (f"Recall is {tn / (tn + fp):.1%} on negative reviews and {tp / (tp + fn):.1%} on positive reviews: "
          f"{fp:,} negative reviews are predicted positive and {fn:,} positive reviews are predicted negative. "
          f"{TODO.replace('TO FILL AFTER RUN', 'TO FILL AFTER RUN: inspect a few misclassified reviews (sarcasm, mixed side-effects vs efficacy) and comment')}")
    short = f"Final: **{best['name']}**: test acc {t['accuracy']:.3f}, macro-F1 {t['macro_f1']:.3f}"
    return s, cm, short


def fill(template, values):
    for k, v in values.items():
        template = template.replace("{{" + k + "}}", str(v))
    return template


def drop_missing_images(md, base_dir):
    def repl(m):
        path = os.path.normpath(os.path.join(base_dir, m.group(2)))
        return m.group(0) if os.path.exists(path) else f"{TODO} *(figure {os.path.basename(path)})*"
    return re.sub(r"!\[([^\]]*)\]\(([^)]+)\)(\{[^}]*\})?", repl, md)


# ---------- LaTeX report (report/report.tex + generated report/results.tex) ----------
def tex_inline(s):
    """Markdown-ish inline text -> LaTeX (escape specials, **bold**, *italic*, `code`, TODO markers)."""
    s = str(s)
    for a, b in (("\\", r"\textbackslash{}"), ("&", r"\&"), ("%", r"\%"), ("#", r"\#"), ("_", r"\_"),
                 ("$", r"\$"), ("~", r"\textasciitilde{}"), ("^", r"\^{}"), ("\u2212", "-"), ("\u2265", r"$\geq$"),
                 ("\u2264", r"$\leq$"), ("\u2248", r"$\approx$")):
        s = s.replace(a, b)
    s = re.sub(r"\*\*\[TO FILL AFTER RUN(.*?)\]\*\*", lambda m: r"\todo{" + m.group(1) + "}", s)
    s = re.sub(r"\*\*(.+?)\*\*", r"\\textbf{\1}", s)
    s = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\\emph{\1}", s)
    s = re.sub(r"`(.+?)`", r"\\texttt{\1}", s)
    return s


def tex_table(md):
    rows = [l for l in md.strip().splitlines() if l.startswith("|") and not re.match(r"^\|[-| :]+\|$", l)]
    if not rows:
        return tex_inline(md)
    cells = [[tex_inline(c.strip()) for c in r.strip().strip("|").split("|")] for r in rows]
    n = len(cells[0])
    body = [r"\begin{tabular}{l" + "r" * (n - 1) + "}", r"\toprule",
            " & ".join(r"\textbf{%s}" % c for c in cells[0]) + r" \\", r"\midrule"]
    body += [" & ".join(r) + r" \\" for r in cells[1:]]
    body += [r"\bottomrule", r"\end{tabular}"]
    tab = "\n".join(body)
    return r"\resizebox{\linewidth}{!}{%" + "\n" + tab + "}" if n > 5 else tab


def write_latex(values, res, out_dir="report", fig_dir="../figures/"):
    abstract = TODO
    if res:
        best = max(res.values(), key=lambda r: r["best_val_macro_f1"])
        abstract = (f"The final model, selected on validation macro-F1, is **{best['name']}**, which reaches a test "
                    f"accuracy of {best['test']['accuracy']:.3f} and a test macro-F1 of {best['test']['macro_f1']:.3f}.")
    macros = {
        "NTrainRaw": "N_TRAIN_RAW", "NTestRaw": "N_TEST_RAW", "NTrainKept": "N_TRAIN_KEPT", "NTestKept": "N_TEST_KEPT",
        "Dropped": "DROPPED", "PosTrain": "POS_TRAIN", "PosTest": "POS_TEST", "LenMean": "LEN_MEAN",
        "LenMedian": "LEN_MEDIAN", "LenPNF": "LEN_P95", "Coverage": "COVERAGE", "NCond": "N_COND",
        "NDrugs": "N_DRUGS", "NTrain": "N_TRAIN", "NVal": "N_VAL", "NTest": "N_TEST", "Vocab": "VOCAB",
        "HparamFindings": "HPARAM_FINDINGS", "ArchFindings": "ARCH_FINDINGS", "ProjFindings": "PROJ_FINDINGS",
        "BestSummary": "BEST_SUMMARY", "CmSentence": "CM_SENTENCE",
    }
    lines = ["% AUTO-GENERATED by `python -m src.build_docs` from results/ -- do not edit by hand.",
             r"\newcommand{\figdir}{%s}" % fig_dir]
    lines += [r"\newcommand{\%s}{%s}" % (m, tex_inline(values[k])) for m, k in macros.items()]
    lines.append(r"\newcommand{\AbstractResult}{%s}" % tex_inline(abstract))
    for m, k in (("ParamsTable", "PARAMS_TABLE"), ("ArchTable", "ARCH_TABLE"), ("HparamTable", "HPARAM_TABLE")):
        lines.append(r"\newcommand{\%s}{%%" % m + "\n" + tex_table(values[k]) + "}")
    path = os.path.join(out_dir, "results.tex")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    return path


def overleaf_bundle(values, res, out_zip="report/overleaf.zip"):
    """report.tex + results.tex (figures/ path) + figures, ready to upload to Overleaf."""
    import tempfile
    import zipfile
    with tempfile.TemporaryDirectory() as tmp:
        write_latex(values, res, out_dir=tmp, fig_dir="figures/")
        with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_DEFLATED) as z:
            z.write("report/report.tex", "report.tex")
            z.write(os.path.join(tmp, "results.tex"), "results.tex")
            for p in sorted(glob.glob("figures/*.png")):
                z.write(p, "figures/" + os.path.basename(p))
    return out_zip


# ---------- markdown -> docx fallback ----------
def md_to_docx(md_path, out_path):
    import docx
    from docx.shared import Inches, Pt

    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.8)
    base = os.path.dirname(md_path)

    def add_runs(par, text):
        for i, chunk in enumerate(re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)):
            if not chunk:
                continue
            if chunk.startswith("**"):
                par.add_run(chunk[2:-2]).bold = True
            elif chunk.startswith("`"):
                r = par.add_run(chunk[1:-1])
                r.font.name = "Consolas"
            elif chunk.startswith("*") and len(chunk) > 2:
                par.add_run(chunk[1:-1]).italic = True
            else:
                par.add_run(chunk)

    lines = open(md_path, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                p = doc.add_paragraph()
                r = p.add_run(lines[i])
                r.font.name, r.font.size = "Consolas", Pt(9)
                p.paragraph_format.space_after = Pt(0)
                i += 1
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(line.lstrip("#").strip(), level=min(level, 3) - 1 if level > 1 else 0)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[-| :]+\|$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Light Grid Accent 1"
            for r, cells in enumerate(rows):
                for c, txt in enumerate(cells):
                    cell = t.cell(r, c)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], txt)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(8.5)
            continue
        elif re.match(r"^!\[", line):
            imgs = re.findall(r"!\[[^\]]*\]\(([^)]+)\)(?:\{width=(\d+)%\})?", line)
            p = doc.add_paragraph()
            for path, w in imgs:
                width = 6.8 * (int(w) if w else 90) / 100
                p.add_run().add_picture(os.path.join(base, path), width=Inches(width))
            rest = re.sub(r"!\[[^\]]*\]\([^)]+\)(\{[^}]*\})?", "", line).strip()
            if rest:
                add_runs(doc.add_paragraph(), rest)
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif line.strip():
            add_runs(doc.add_paragraph(), line)
        i += 1
    doc.save(out_path)


def export(report_md, slides_md):
    out = []
    if shutil.which("pandoc"):
        pdf = report_md.replace(".md", ".pdf")
        r = subprocess.run(["pandoc", os.path.basename(report_md), "-o", os.path.basename(pdf),
                            "-V", "geometry:margin=2cm"], cwd=os.path.dirname(report_md))
        if r.returncode == 0:
            out.append(pdf)
    if not out:
        docx_path = report_md.replace(".md", ".docx")
        md_to_docx(report_md, docx_path)
        out.append(docx_path)
    marp = shutil.which("marp")
    if marp:
        for flag, ext in (("--pdf", ".pdf"), ("--pptx", ".pptx")):
            r = subprocess.run([marp, slides_md, flag, "--allow-local-files", "-o", slides_md.replace(".md", ext)])
            if r.returncode == 0:
                out.append(slides_md.replace(".md", ext))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default="results")
    ap.add_argument("--team", default="Mohamed Amine Lamrid, Badr Misraoui (Lab group E)")
    ap.add_argument("--repo-url", default="https://github.com/amine2003lamrid-hue/drug-review-lstm")
    args = ap.parse_args()

    with open("configs/base.yaml") as f:
        cfg = yaml.safe_load(f)
    res = load(args.results)
    eda = {}
    if os.path.exists(f"{args.results}/eda_stats.json"):
        with open(f"{args.results}/eda_stats.json") as f:
            eda = json.load(f)
    vocab_size = None
    try:
        from .run_experiments import get_data
        vocab_size = len(get_data(cfg)[0])
    except Exception as e:  # data not downloaded: leave placeholders
        print("vocab unavailable:", e)

    g = lambda k, f=str: f(eda[k]) if k in eda else TODO  # noqa: E731
    arch_long, _ = arch_findings(res)
    hp_long, hp_short = hparam_findings(res)
    best_long, cm_sentence, best_short = best_summary(res)
    values = {
        "TEAM": args.team, "REPO_URL": args.repo_url,
        "N_TRAIN_RAW": g("n_train_raw", lambda v: f"{v:,}"), "N_TEST_RAW": g("n_test_raw", lambda v: f"{v:,}"),
        "N_TRAIN_KEPT": g("n_train_kept", lambda v: f"{v:,}"), "N_TEST_KEPT": g("n_test_kept", lambda v: f"{v:,}"),
        "DROPPED": g("dropped_neutral", lambda v: f"{v:,}"),
        "POS_TRAIN": g("pos_ratio_train", pct), "POS_TEST": g("pos_ratio_test", pct),
        "LEN_MEAN": g("len_mean", lambda v: f"{v:.0f}"), "LEN_MEDIAN": g("len_median", lambda v: f"{v:.0f}"),
        "LEN_P95": g("len_p95", lambda v: f"{v:.0f}"), "COVERAGE": g("coverage_max_len", pct),
        "N_COND": g("n_conditions", lambda v: f"{v:,}"), "N_DRUGS": g("n_drugs", lambda v: f"{v:,}"),
        "N_TRAIN": f"{cfg['n_train']:,}", "N_VAL": f"{cfg['n_val']:,}", "N_TEST": f"{cfg['n_test']:,}",
        "VOCAB": f"{vocab_size:,}" if vocab_size else TODO,
        "PARAMS_TABLE": params_table(cfg, vocab_size),
        "ARCH_TABLE": table_from(f"{args.results}/architecture_table.md"),
        "ARCH_TABLE_SHORT": table_from(f"{args.results}/architecture_table.md", short=True),
        "HPARAM_TABLE": table_from(f"{args.results}/hparam_table.md"),
        "ARCH_FINDINGS": arch_long, "HPARAM_FINDINGS": hp_long, "HPARAM_FINDINGS_SHORT": hp_short,
        "PROJ_FINDINGS": proj_findings(res),
        "BEST_SUMMARY": best_long, "BEST_SUMMARY_SHORT": best_short, "CM_SENTENCE": cm_sentence,
    }
    for tpl, out in (("report/report_template.md", "report/report.md"), ("slides/slides_template.md", "slides/slides.md")):
        with open(tpl, encoding="utf-8") as f:
            text = fill(f.read(), values)
        text = drop_missing_images(text, os.path.dirname(out))
        with open(out, "w", encoding="utf-8") as f:
            f.write(text)
    if os.path.exists("README.md") and os.path.exists(f"{args.results}/architecture_table.md"):
        with open("README.md", encoding="utf-8") as f:
            readme = f.read()
        block = "\n".join([
            "Architecture comparison (test set):", "", values["ARCH_TABLE"], "",
            "Hyperparameter study (vanilla LSTM, test set):", "", values["HPARAM_TABLE"], "",
            best_long, "", "Full discussion: `report/report.md`; figures in `figures/`.",
        ])
        readme = re.sub(r"<!-- RESULTS_START -->.*<!-- RESULTS_END -->",
                        lambda _: f"<!-- RESULTS_START -->\n{block}\n<!-- RESULTS_END -->", readme, flags=re.S)
        with open("README.md", "w", encoding="utf-8") as f:
            f.write(readme)
    left = sum(open(p, encoding="utf-8").read().count("TO FILL AFTER RUN") for p in ("report/report.md", "slides/slides.md"))
    print("written report/report.md, slides/slides.md;", left, "placeholders left")
    tex = write_latex(values, res)
    print("written", tex, "and", overleaf_bundle(values, res))
    if shutil.which("pdflatex"):
        for _ in range(2):
            subprocess.run(["pdflatex", "-interaction=nonstopmode", "report.tex"], cwd="report",
                           stdout=subprocess.DEVNULL)
    print("exported:", export("report/report.md", "slides/slides.md"))


if __name__ == "__main__":
    main()
